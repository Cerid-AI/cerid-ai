# Copyright (c) 2026 Cerid AI. All rights reserved.
# SPDX-License-Identifier: FSL-1.1-ALv2

"""The layout-aware parser never feeds raw MIME or binary into the hash,
metadata extraction or categorisation.

``layout_aware_parse`` used to return the file's bytes decoded as text as the
canonical text for every format it claims. For PDF, DOCX, XLSX and ``.eml``
that text is the container (``%PDF``, zip members, MIME headers and base64
attachments), so the content hash, the NLP metadata and the AI category were
all computed on garbage while the chunks came from the real parser. The
canonical text for those formats is now the parser's extracted text, and the
hash is defined on it. Existing artifacts carry the old hash, so dedup
consults both.
"""
from __future__ import annotations

import hashlib
import re
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate

import app.services.ingestion as ing
import config
from app.services import folder_scanner
from app.services.ingestion import _content_hash, ingest_content
from core.ingest.dispatch import layout_aware_parse, legacy_hash_text

BODY = "Revenue rose in the third quarter and costs were flat."
_BASE64_RUN = re.compile(r"[A-Za-z0-9+/]{40,}={0,2}")


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _write_pdf(path: Path) -> Path:
    SimpleDocTemplate(str(path), pagesize=letter).build(
        [Paragraph(BODY, getSampleStyleSheet()["Normal"])],
    )
    return path


def _write_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Quarterly figures", level=1)
    doc.add_paragraph(BODY)
    doc.save(str(path))
    return path


def _write_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["quarter", "note"])
    ws.append(["Q3", BODY])
    wb.save(str(path))
    return path


def _write_eml(path: Path) -> Path:
    msg = EmailMessage()
    msg["From"] = "Sender <sender@example.test>"
    msg["To"] = "Reader <reader@example.test>"
    msg["Subject"] = "Quarterly figures"
    msg["Message-ID"] = "<figures-0001@example.test>"
    msg["Date"] = "Mon, 07 Sep 2026 09:30:00 +0000"
    msg.set_content(BODY + "\n")
    msg.add_attachment(
        bytes(range(256)) * 8, maintype="application", subtype="octet-stream",
        filename="figures.bin",
    )
    path.write_bytes(msg.as_bytes())
    return path


_WRITERS = {"pdf": _write_pdf, "docx": _write_docx, "xlsx": _write_xlsx, "eml": _write_eml}


@pytest.fixture(params=sorted(_WRITERS))
def binary_file(request, tmp_path: Path) -> Path:
    return _WRITERS[request.param](tmp_path / f"figures.{request.param}")


def _assert_is_extracted_text(text: str) -> None:
    assert BODY in text
    assert "Content-Type:" not in text
    assert "Content-Transfer-Encoding" not in text
    assert "\x00" not in text
    assert "�" not in text
    assert not text.startswith(("%PDF", "PK"))
    assert _BASE64_RUN.search(text) is None, "base64 run survived into the canonical text"


# ---------------------------------------------------------------------------
# dispatch
# ---------------------------------------------------------------------------


def test_canonical_text_is_the_parsers_text_for_binary_and_mime_formats(binary_file: Path):
    result = layout_aware_parse(binary_file)
    assert result is not None
    text, chunks = result

    _assert_is_extracted_text(text)
    assert chunks
    assert text != binary_file.read_text(encoding="utf-8", errors="replace")


def test_plain_text_formats_keep_the_file_text(tmp_path: Path):
    md = tmp_path / "note.md"
    md.write_text("# Top\n## Sub\nLeaf body text.\n", encoding="utf-8")

    text, _ = layout_aware_parse(md)

    assert text == md.read_text(encoding="utf-8")
    assert legacy_hash_text(md) is None


def test_legacy_hash_text_is_the_bytes_decoded(binary_file: Path):
    assert legacy_hash_text(binary_file) == binary_file.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------------------
# ingest_file: hash, metadata and categorisation see the extracted text
# ---------------------------------------------------------------------------


class _Stores:
    def __init__(self) -> None:
        self.writes: list[dict[str, Any]] = []
        self.categorised: list[str] = []

    def ingest_content(
        self, content: str, domain: str = "general",
        metadata: dict[str, Any] | None = None, **kwargs: Any,
    ) -> dict[str, Any]:
        self.writes.append({"content": content, "metadata": dict(metadata or {}), **kwargs})
        return {"status": "success", "artifact_id": "art-1", "domain": domain, "chunks": 1}

    async def ai_categorize(self, text: str, filename: str, mode: str) -> dict[str, Any]:
        self.categorised.append(text)
        return {}


@pytest.fixture
def stores(monkeypatch, tmp_path: Path) -> _Stores:
    fake = _Stores()
    monkeypatch.setattr(ing, "ingest_content", fake.ingest_content)
    monkeypatch.setattr(ing, "ai_categorize", fake.ai_categorize)
    monkeypatch.setattr(ing.graph, "write_has_attachment", lambda **_: True)
    monkeypatch.setattr(ing, "get_neo4j", lambda: object())
    monkeypatch.setattr(config, "ARCHIVE_PATH", str(tmp_path))
    monkeypatch.setattr(config, "ENABLE_LAYOUT_AWARE_PARSING", True)
    return fake


@pytest.mark.asyncio
async def test_ingest_file_hashes_and_categorises_the_extracted_text(stores: _Stores, binary_file: Path):
    extracted, _ = layout_aware_parse(binary_file)

    await ing.ingest_file(str(binary_file), domain="", categorize_mode="ai", skip_metadata=True)

    write = stores.writes[0]
    assert write["content"] == extracted
    assert _content_hash(write["content"]) == _sha(extracted)
    assert write["prior_hashes"] == (_sha(legacy_hash_text(binary_file)),)
    assert write["metadata"]["content_hash_version"] == "2"
    assert write["pre_chunked"]
    _assert_is_extracted_text(stores.categorised[0])
    for value in write["metadata"].values():
        if isinstance(value, str):
            assert "Content-Type:" not in value and "\x00" not in value


@pytest.mark.asyncio
async def test_plain_text_ingest_carries_no_prior_hash(stores: _Stores, tmp_path: Path):
    md = tmp_path / "note.md"
    md.write_text("# Top\n\nLeaf body text.\n", encoding="utf-8")

    await ing.ingest_file(str(md), domain="general", skip_metadata=True)

    write = stores.writes[0]
    assert write["prior_hashes"] == ()
    assert "content_hash_version" not in write["metadata"]


# ---------------------------------------------------------------------------
# dedup consults both hash versions
# ---------------------------------------------------------------------------


def _neo4j_holding(content_hash: str) -> MagicMock:
    """A driver whose only artifact carries ``content_hash``; the dedup query
    matches when that hash is among the ones asked for."""
    session = MagicMock()

    def run(query: str, **params: Any) -> MagicMock:
        result = MagicMock()
        hashes = params.get("hashes") or []
        hit = {"id": "old-pdf", "filename": "figures.pdf", "domain": "general"}
        result.single.return_value = hit if content_hash in hashes else None
        return result

    session.run.side_effect = run
    driver = MagicMock()
    driver.session.return_value.__enter__ = MagicMock(return_value=session)
    driver.session.return_value.__exit__ = MagicMock(return_value=False)
    return driver


def test_reingest_of_a_file_hashed_by_an_earlier_release_is_a_duplicate():
    legacy = _sha("%PDF-1.4 �� garbage")
    with patch("app.services.ingestion.get_redis", return_value=MagicMock()), \
         patch("app.services.ingestion.get_chroma", return_value=MagicMock()), \
         patch("app.services.ingestion.get_neo4j", return_value=_neo4j_holding(legacy)):
        result = ingest_content(
            BODY, domain="general", metadata={"filename": "figures.pdf"}, prior_hashes=(legacy,),
        )

    assert result["status"] == "duplicate"
    assert result["artifact_id"] == "old-pdf"


def test_reingest_of_the_same_extracted_text_is_a_duplicate():
    with patch("app.services.ingestion.get_redis", return_value=MagicMock()), \
         patch("app.services.ingestion.get_chroma", return_value=MagicMock()), \
         patch("app.services.ingestion.get_neo4j", return_value=_neo4j_holding(_sha(BODY))):
        result = ingest_content(BODY, domain="general", metadata={"filename": "figures.pdf"})

    assert result["status"] == "duplicate"


def test_watched_folder_links_an_artifact_hashed_by_an_earlier_release(monkeypatch, tmp_path: Path):
    pdf = _write_pdf(tmp_path / "figures.pdf")
    legacy = _sha(legacy_hash_text(pdf))
    stamped: list[tuple[str, str, str]] = []
    monkeypatch.setattr(config, "ENABLE_LAYOUT_AWARE_PARSING", True)
    monkeypatch.setattr(folder_scanner, "get_neo4j", lambda: _neo4j_holding(legacy))
    monkeypatch.setattr(folder_scanner, "_stamp_artifact", lambda *a: stamped.append(a))

    assert folder_scanner._link_file(str(pdf), "folder-1") is True
    assert stamped == [("old-pdf", "general", "folder-1")]
